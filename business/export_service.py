"""
엑셀 내보내기 서비스
"""

from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database.db_manager import DatabaseManager


class ExportService:
    """엑셀 내보내기 서비스"""

    def __init__(self):
        self.db = DatabaseManager()

    # ============ 스타일 정의 ============

    @staticmethod
    def get_header_style():
        """헤더 셀 스타일"""
        return {
            'font': Font(bold=True, color="FFFFFF"),
            'fill': PatternFill(start_color="366092", end_color="366092", fill_type="solid"),
            'alignment': Alignment(horizontal="center", vertical="center", wrap_text=True),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }

    @staticmethod
    def get_data_style():
        """데이터 셀 스타일"""
        return {
            'alignment': Alignment(horizontal="left", vertical="center"),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }

    @staticmethod
    def get_center_style():
        """중앙 정렬 셀 스타일"""
        return {
            'alignment': Alignment(horizontal="center", vertical="center"),
            'border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )
        }

    # ============ 헤더 작성 ============

    def _write_header(self, sheet, headers):
        """시트에 헤더 작성"""
        style = self.get_header_style()
        for col, header in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col, value=header)
            cell.font = style['font']
            cell.fill = style['fill']
            cell.alignment = style['alignment']
            cell.border = style['border']

    # ============ 자원봉사자 목록 내보내기 ============

    def export_all_volunteers(self, filepath):
        """모든 자원봉사자 목록 내보내기"""
        volunteers = self.db.get_all_volunteers()

        wb = Workbook()
        ws = wb.active
        ws.title = "자원봉사자 목록"

        headers = ["이름", "성별", "생년월일", "소속", "핸드폰번호", "PIN", "등록일", "활성상태"]
        self._write_header(ws, headers)

        data_style = self.get_data_style()
        center_style = self.get_center_style()

        for row, volunteer in enumerate(volunteers, 2):
            ws.cell(row=row, column=1, value=volunteer['name'])
            ws.cell(row=row, column=2, value="남" if volunteer['gender'] == 'M' else "여")
            ws.cell(row=row, column=3, value=volunteer['birth_date'])
            ws.cell(row=row, column=4, value=volunteer['organization'])
            ws.cell(row=row, column=5, value=volunteer['phone_number'])
            ws.cell(row=row, column=6, value=volunteer['pin_number'])
            ws.cell(row=row, column=7, value=volunteer['registered_at'][:10])
            ws.cell(row=row, column=8, value="활성" if volunteer['is_active'] else "비활성")

            for col in range(1, 9):
                cell = ws.cell(row=row, column=col)
                if col == 2 or col == 8:
                    cell.alignment = center_style['alignment']
                else:
                    cell.alignment = data_style['alignment']
                cell.border = data_style['border']

        # 열 너비 조정
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 12

        wb.save(filepath)
        return True

    # ============ 출퇴근 기록 내보내기 ============

    def export_attendance_records(self, filepath, start_date=None, end_date=None):
        """출퇴근 기록 내보내기"""
        records = self.db.get_attendance_records(start_date=start_date, end_date=end_date)

        wb = Workbook()
        ws = wb.active
        ws.title = "출퇴근 기록"

        headers = ["자원봉사자", "봉사처", "출근시간", "퇴근시간", "봉사시간(분)", "날짜"]
        self._write_header(ws, headers)

        data_style = self.get_data_style()

        for row, record in enumerate(records, 2):
            # 자원봉사자 이름
            volunteer = self.db.get_volunteer_by_id(record['volunteer_id'])
            ws.cell(row=row, column=1, value=volunteer['name'])

            # 봉사처 이름
            location = self.db.get_location_by_id(record['location_id'])
            ws.cell(row=row, column=2, value=location['name'])

            # 출근시간
            check_in = record['check_in_time']
            ws.cell(row=row, column=3, value=check_in[11:19] if check_in else "-")

            # 퇴근시간
            check_out = record['check_out_time']
            ws.cell(row=row, column=4, value=check_out[11:19] if check_out else "-")

            # 봉사시간 (분)
            duration = record['duration_minutes']
            ws.cell(row=row, column=5, value=duration if duration else "-")

            # 날짜
            ws.cell(row=row, column=6, value=check_in[:10] if check_in else "-")

            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.alignment = data_style['alignment']
                cell.border = data_style['border']

        # 열 너비 조정
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12

        wb.save(filepath)
        return True

    # ============ 통합 리포트 내보내기 ============

    def export_comprehensive_report(self, filepath):
        """통합 리포트 내보내기 (모든 내용)"""
        wb = Workbook()

        # 시트 1: 자원봉사자 목록
        self._create_volunteers_sheet(wb)

        # 시트 2: 출퇴근 기록
        self._create_attendance_sheet(wb)

        # 시트 3: 월별 통계 (현재 월)
        self._create_monthly_stats_sheet(wb)

        # 시트 4: 자원봉사자별 통계
        self._create_volunteer_stats_sheet(wb)

        # 기본 시트 제거
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])

        wb.save(filepath)
        return True

    def _create_volunteers_sheet(self, wb):
        """자원봉사자 목록 시트 생성"""
        volunteers = self.db.get_all_volunteers()

        ws = wb.create_sheet("자원봉사자 목록")

        headers = ["이름", "성별", "생년월일", "소속", "핸드폰번호", "PIN", "등록일", "활성상태"]
        self._write_header(ws, headers)

        data_style = self.get_data_style()
        center_style = self.get_center_style()

        for row, volunteer in enumerate(volunteers, 2):
            ws.cell(row=row, column=1, value=volunteer['name'])
            ws.cell(row=row, column=2, value="남" if volunteer['gender'] == 'M' else "여")
            ws.cell(row=row, column=3, value=volunteer['birth_date'])
            ws.cell(row=row, column=4, value=volunteer['organization'])
            ws.cell(row=row, column=5, value=volunteer['phone_number'])
            ws.cell(row=row, column=6, value=volunteer['pin_number'])
            ws.cell(row=row, column=7, value=volunteer['registered_at'][:10])
            ws.cell(row=row, column=8, value="활성" if volunteer['is_active'] else "비활성")

            for col in range(1, 9):
                cell = ws.cell(row=row, column=col)
                if col == 2 or col == 8:
                    cell.alignment = center_style['alignment']
                else:
                    cell.alignment = data_style['alignment']
                cell.border = data_style['border']

        # 열 너비 조정
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 10
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 15
        ws.column_dimensions['H'].width = 12

    def _create_attendance_sheet(self, wb):
        """출퇴근 기록 시트 생성"""
        records = self.db.get_attendance_records()

        ws = wb.create_sheet("출퇴근 기록")

        headers = ["자원봉사자", "봉사처", "출근시간", "퇴근시간", "봉사시간(분)", "날짜"]
        self._write_header(ws, headers)

        data_style = self.get_data_style()

        for row, record in enumerate(records, 2):
            # 자원봉사자 이름
            volunteer = self.db.get_volunteer_by_id(record['volunteer_id'])
            ws.cell(row=row, column=1, value=volunteer['name'])

            # 봉사처 이름
            location = self.db.get_location_by_id(record['location_id'])
            ws.cell(row=row, column=2, value=location['name'])

            # 출근시간
            check_in = record['check_in_time']
            ws.cell(row=row, column=3, value=check_in[11:19] if check_in else "-")

            # 퇴근시간
            check_out = record['check_out_time']
            ws.cell(row=row, column=4, value=check_out[11:19] if check_out else "-")

            # 봉사시간 (분)
            duration = record['duration_minutes']
            ws.cell(row=row, column=5, value=duration if duration else "-")

            # 날짜
            ws.cell(row=row, column=6, value=check_in[:10] if check_in else "-")

            for col in range(1, 7):
                cell = ws.cell(row=row, column=col)
                cell.alignment = data_style['alignment']
                cell.border = data_style['border']

        # 열 너비 조정
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 15
        ws.column_dimensions['F'].width = 12

    def _create_monthly_stats_sheet(self, wb):
        """월별 통계 시트 생성"""
        today = datetime.now()
        year = today.year
        month = today.month

        ws = wb.create_sheet("월별 통계")

        # 헤더
        headers = [f"{year}년 {month}월 통계", "자원봉사자", "총 봉사시간", "봉사 횟수"]
        self._write_header(ws, headers)

        # 통계 조회
        from business.attendance_service import AttendanceService
        service = AttendanceService()
        stats = service.get_monthly_statistics(year, month)

        data_style = self.get_data_style()

        row = 2
        for volunteer_id, stat in stats.items():
            ws.cell(row=row, column=2, value=stat['name'])

            # 시간 계산
            hours = stat['total_minutes'] // 60
            minutes = stat['total_minutes'] % 60
            ws.cell(row=row, column=3, value=f"{hours}시간 {minutes}분")

            ws.cell(row=row, column=4, value=stat['visit_count'])

            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.alignment = data_style['alignment']
                cell.border = data_style['border']

            row += 1

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 12

    def _create_volunteer_stats_sheet(self, wb):
        """자원봉사자별 통계 시트 생성"""
        ws = wb.create_sheet("자원봉사자별 통계")

        headers = ["자원봉사자", "소속", "총 봉사시간", "최근 봉사일", "봉사 횟수"]
        self._write_header(ws, headers)

        volunteers = self.db.get_all_volunteers()
        data_style = self.get_data_style()

        row = 2
        for volunteer in volunteers:
            if not volunteer['is_active']:
                continue

            ws.cell(row=row, column=1, value=volunteer['name'])
            ws.cell(row=row, column=2, value=volunteer['organization'])

            # 통계 정보
            stats = self.db.get_volunteer_stats(volunteer['id'])

            # 시간 계산
            if stats and stats['total_minutes']:
                hours = stats['total_minutes'] // 60
                minutes = stats['total_minutes'] % 60
                ws.cell(row=row, column=3, value=f"{hours}시간 {minutes}분")
            else:
                ws.cell(row=row, column=3, value="0시간")

            # 최근 봉사일
            if stats and stats['last_visit']:
                ws.cell(row=row, column=4, value=stats['last_visit'][:10])
            else:
                ws.cell(row=row, column=4, value="-")

            # 봉사 횟수
            ws.cell(row=row, column=5, value=stats['visit_count'] if stats else 0)

            for col in range(1, 6):
                cell = ws.cell(row=row, column=col)
                cell.alignment = data_style['alignment']
                cell.border = data_style['border']

            row += 1

        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 12
